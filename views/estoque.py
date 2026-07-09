import flet as ft
from datetime import datetime
from openpyxl import Workbook

from utils.export_manager import (
    novo_pdf,
    novo_excel,
    abrir_pdf,
    abrir_excel
)

from database.estoque_supabase import (
    listar_estoque,
    atualizar_produto,
    excluir_lote,
    registrar_ajuste_estoque
)

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)

from views.tema import (
    AZUL,
    VERDE,
    BRANCO,
    FUNDO,
    TEXTO,
    CINZA,
    PADDING_CARD,
    ESPACAMENTO,
    card_padrao,
    botao_primario,
    botao_sucesso,
    botao_secundario,
    titulo_pagina,
    container_pagina
)

from app_paths import LOGO_PATHS
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from utils.file_manager import salvar_arquivo

from pathlib import Path

from openpyxl.drawing.image import Image as ExcelImage

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm

# REPORTS_DIR = BASE_DIR / "reports"

# LOGO_PATHS = [

#    BASE_DIR / "assets" / "logo.png",

#    BASE_DIR / "assets" / "logo" / "logo_nova.png",

#    BASE_DIR / "logo.png"

#]

from pathlib import Path

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    Image
)

from reportlab.lib.pagesizes import A4, landscape

from reportlab.lib.units import cm

from openpyxl.drawing.image import Image as ExcelImage

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

from openpyxl.utils import get_column_letter


def estoque_view(page, usuario):

    largura_tela = page.width or 390
    altura_tela = page.height or 700

    largura_util = max(largura_tela - 30, 300)

    largura_campo_medio = min(largura_util, 300)
    largura_campo_grande = min(largura_util, 400)

    largura_card = min(largura_util, 180)

    largura_dialog = min(largura_util, 700)
    altura_dialog = min(max(altura_tela - 180, 350), 500)

    estoque = listar_estoque()

    def atualizar_estoque():

        nonlocal estoque

        estoque = listar_estoque()

    detalhes_produto = ft.Column()

    editar_nome = ft.TextField(label="Nome")

    editar_unidade = ft.TextField(
        label="Unidade"
    )

    editar_quantidade = ft.TextField(
        label="Quantidade"
    )

    editar_lote = ft.TextField(
        label="Lote"
    )

    editar_validade = ft.TextField(
        label="Validade"
    )

    produto_editando_id = None
    lote_ids_excluindo = []
    lote_editando_id = None
    produto_original = None

    lote_ids_editando = []
    lote_ids_excluindo = []

    dialog_produto = ft.AlertDialog(

        modal=True,

        title=ft.Text(
            "Detalhes do Produto"
        ),

        content=ft.Container(
            content=detalhes_produto,
            width=min(largura_util, 500),
            height=min(altura_dialog, 350)
        ),

        actions=[

            botao_secundario(
                "Fechar",
                icon=ft.Icons.CLOSE,
                on_click=lambda e: fechar_dialog(e)
            )

        ]

    )

    def salvar_edicao(e):

        nonlocal produto_editando_id
        nonlocal produto_original

        antes = (
            f"Nome: {produto_original['nome']} | "
            f"Unidade: {produto_original['unidade']} | "
            f"Quantidade: {produto_original['quantidade']} | "
            f"Lote: {produto_original['lote']} | "
            f"Validade: {produto_original['validade']}"
        )

        depois = (
            f"Nome: {editar_nome.value} | "
            f"Unidade: {editar_unidade.value} | "
            f"Quantidade: {editar_quantidade.value} | "
            f"Lote: {editar_lote.value} | "
            f"Validade: {editar_validade.value}"
        )

        usuario_nome = (
            usuario.get("nome") or
            usuario.get("email") or
            usuario.get("login") or
            "Usuário não identificado"
        )

        atualizar_produto(

            produto_editando_id,
            lote_ids_editando,

            editar_nome.value,

            editar_unidade.value,

            int(editar_quantidade.value),

            editar_lote.value,

            editar_validade.value

        )

        registrar_ajuste_estoque(
            produto_editando_id,
            lote_ids_editando[0],
            editar_nome.value,
            usuario_nome,
            "Edição de estoque",
            antes,
            depois
        )

        atualizar_estoque()
        atualizar_cards()

        carregar_tabela(
            campo_pesquisa.value,
            filtro_dropdown.value
        )

        dialog_editar.open = False

        e.page.snack_bar = ft.SnackBar(
            content=ft.Text(
                "Produto atualizado com sucesso!"
            )
        )

        e.page.snack_bar.open = True

        e.page.update()

    def fechar_edicao(e):

        dialog_editar.open = False

        e.page.update()

    def confirmar_exclusao(e):

        nonlocal lote_ids_excluindo

        excluir_lote(lote_ids_excluindo)

        atualizar_estoque()
        atualizar_cards()

        carregar_tabela(
            campo_pesquisa.value,
            filtro_dropdown.value
        )

        dialog_excluir.open = False

        e.page.snack_bar = ft.SnackBar(
            content=ft.Text(
                "Lote excluído com sucesso!"
            )
        )

        e.page.snack_bar.open = True

        e.page.update()        

    dialog_editar = ft.AlertDialog(

        modal=True,

        title=ft.Text("Editar Produto"),

        content=ft.Container(
            content=ft.Column(
                [
                    editar_nome,
                    editar_unidade,
                    editar_quantidade,
                    editar_lote,
                    editar_validade
                ]
            ),
            width=min(largura_util, 500)
        ),

        actions=[

            botao_sucesso(
                "Salvar",
                icon=ft.Icons.SAVE,
                on_click=salvar_edicao
            ),

            botao_secundario(
                "Cancelar",
                icon=ft.Icons.CLOSE,
                on_click=fechar_edicao
            )

        ]

    )

    dialog_excluir = ft.AlertDialog(

        modal=True,

        title=ft.Text(
            "Excluir Produto"
        ),

        content=ft.Text(
            "Deseja realmente excluir este produto?"
        ),

        actions=[

            botao_secundario(
                "Cancelar",
                icon=ft.Icons.CLOSE,
                on_click=lambda e: (
                    setattr(dialog_excluir, "open", False),
                    e.page.update()
                )
            ),

            botao_secundario(
                "Excluir",
                icon=ft.Icons.DELETE,
                on_click=confirmar_exclusao
            )

        ]

    )

    def fechar_dialog(e):

        dialog_produto.open = False

        e.page.update()             

    def visualizar_produto(e, item):

        detalhes_produto.controls.clear()

        detalhes_produto.controls.extend(

            [

                ft.Text(
                    f"Produto: {item['nome']}"
                ),

                ft.Text(
                    f"Unidade: {item['unidade_medida']}"
                ),

                ft.Text(
                    f"Quantidade: {item['quantidade']}"
                ),

                ft.Text(
                    f"Lote: {item['lote']}"
                ),

                ft.Text(
                    f"Validade: {item['validade']}"
                )

            ]

        )

        if dialog_produto not in e.page.overlay:
            e.page.overlay.append(dialog_produto)

        dialog_produto.open = True

        e.page.update()

    total_itens = len(estoque)

    total_baixo = 0
    total_vencidos = 0

    for item in estoque:

        try:

            validade = datetime.strptime(
                item["validade"],
                "%d/%m/%Y"
            )

            if validade < datetime.now():

                total_vencidos += 1

        except:

            pass

        if item["quantidade"] <= item["estoque_minimo"]:

            total_baixo += 1

    texto_total_itens = ft.Text(
        str(total_itens),
        size=28
    )

    texto_total_baixo = ft.Text(
        str(total_baixo),
        size=28
    )

    texto_total_vencidos = ft.Text(
        str(total_vencidos),
        size=28
    )

    def atualizar_cards():

        total_itens = len(estoque)

        total_baixo = 0
        total_vencidos = 0

        for item in estoque:

            try:
                validade = datetime.strptime(
                    item["validade"],
                    "%d/%m/%Y"
                )

                if validade < datetime.now():
                    total_vencidos += 1

            except:
                pass

            if item["quantidade"] <= item["estoque_minimo"]:
                total_baixo += 1

        texto_total_itens.value = str(total_itens)
        texto_total_baixo.value = str(total_baixo)
        texto_total_vencidos.value = str(total_vencidos)        

    campo_pesquisa = ft.TextField(
        label="Pesquisar Produto",
        prefix_icon=ft.Icons.SEARCH,
        width=largura_campo_grande
    )
    
    filtro_dropdown = ft.Dropdown(
        label="Filtro",
        width=largura_campo_medio,
        value="TODOS",
        options=[
            ft.dropdown.Option("TODOS"),
            ft.dropdown.Option("VENCIDOS"),
            ft.dropdown.Option("ESTOQUE BAIXO")
        ]
    )

    ordenacao_dropdown = ft.Dropdown(
        label="Ordenar por",
        width=largura_campo_medio,
        value="nome",
        options=[
            ft.dropdown.Option("nome"),
            ft.dropdown.Option("quantidade"),
            ft.dropdown.Option("validade")
        ]
    )

    def atualizar_filtros(e):

        nonlocal pagina_atual


        pagina_atual = 1

        atualizar_estoque()
        atualizar_cards()

        carregar_tabela(
            campo_pesquisa.value,
            filtro_dropdown.value
        )

        e.page.update()

    filtro_dropdown.on_select = atualizar_filtros

    ordenacao_dropdown.on_select = lambda e: (
        carregar_tabela(
            campo_pesquisa.value,
            filtro_dropdown.value
        ),
        e.page.update()
    )  

    tabela_estoque = ft.DataTable(

        columns=[

            ft.DataColumn(ft.Text("Status")),
            ft.DataColumn(ft.Text("Produto")),
            ft.DataColumn(ft.Text("Unidade")),
            ft.DataColumn(ft.Text("Quantidade")),
            ft.DataColumn(ft.Text("Lote")),
            ft.DataColumn(ft.Text("Validade")),
            ft.DataColumn(ft.Text("Ações"))
              
        ],

        rows=[]

    )

    itens_por_pagina = 10
    pagina_atual = 1

    texto_pagina = ft.Text(
        "Página 1"
    )

    def editar_produto(e, item):

        nonlocal produto_editando_id
        nonlocal lote_editando_id
        nonlocal lote_ids_editando
        nonlocal produto_original

        produto_editando_id = item["produto_id"]

        lote_editando_id = item["lote_id"]

        lote_ids_editando = item.get(
            "lote_ids",
            [item["lote_id"]]
        )

        produto_original = {
            "nome": item["nome"],
            "unidade": item["unidade_medida"],
            "quantidade": item["quantidade"],
            "lote": item["lote"],
            "validade": item["validade"]
        }

        editar_nome.value = item["nome"]

        editar_unidade.value = item["unidade_medida"]

        editar_quantidade.value = str(
            item["quantidade"]
        )

        editar_lote.value = item["lote"]

        editar_validade.value = item["validade"]

        if dialog_editar not in e.page.overlay:
            e.page.overlay.append(dialog_editar)

        dialog_editar.open = True

        e.page.update()

    def excluir_produto(e, item):
        
        nonlocal lote_ids_excluindo

        lote_ids_excluindo = item.get(
            "lote_ids",
            [item["lote_id"]]
        )

        if dialog_excluir not in e.page.overlay:
            e.page.overlay.append(dialog_excluir)

        dialog_excluir.open = True

        e.page.update()          

    def carregar_tabela(
        filtro="",
        tipo_filtro="TODOS"
    ):

        nonlocal pagina_atual

        tabela_estoque.rows.clear()

        itens_ordenados = []

        for item in estoque:

            if filtro:

                if filtro.lower() not in item["nome"].lower():
                    continue

            status = "🟢"

            try:

                validade = datetime.strptime(
                    item["validade"],
                    "%d/%m/%Y"
                )

                hoje = datetime.now()

                if validade < hoje:

                    status = "🔴"

                elif item["quantidade"] <= item["estoque_minimo"]:

                    status = "🟡"

                if tipo_filtro == "VENCIDOS" and status != "🔴":
                    continue

                if tipo_filtro == "ESTOQUE BAIXO" and status != "🟡":
                    continue    

            except Exception as erro:

                print("ERRO DATA:", erro)

            prioridade = 3

            if status == "🔴":
                prioridade = 1

            elif status == "🟡":
                prioridade = 2

            itens_ordenados.append(
                (
                    prioridade,
                    status,
                    item
                )
            )
            
        if ordenacao_dropdown.value == "nome":
            itens_ordenados.sort(
                key=lambda x: x[2]["nome"].lower()
            )

        elif ordenacao_dropdown.value == "quantidade":
            itens_ordenados.sort(
                key=lambda x: x[2]["quantidade"]
            )

        elif ordenacao_dropdown.value == "validade":

            def data_ordenacao(x):
                try:
                    return datetime.strptime(
                        x[2]["validade"],
                        "%d/%m/%Y"
                    )
                except:
                    return datetime.max

            itens_ordenados.sort(
                key=data_ordenacao
            )

        total_paginas = max(
            1,
            (len(itens_ordenados) + itens_por_pagina - 1) // itens_por_pagina
        )

        if pagina_atual > total_paginas:
            pagina_atual = total_paginas

        inicio = (pagina_atual - 1) * itens_por_pagina
        fim = inicio + itens_por_pagina

        itens_pagina = itens_ordenados[inicio:fim]

        texto_pagina.value = f"Página {pagina_atual} de {total_paginas}"

        botao_anterior.disabled = pagina_atual <= 1
        botao_proximo.disabled = pagina_atual >= total_paginas

        for prioridade, status, item in itens_pagina:

            # cria a lista de ações
            acoes = [

                ft.IconButton(
                    icon=ft.Icons.VISIBILITY,
                    tooltip="Visualizar",
                    on_click=lambda e, p=item: visualizar_produto(e, p)
                )

            ]

            if usuario["perfil"] != "Olheiro":

                acoes.extend([

                    ft.IconButton(
                        icon=ft.Icons.EDIT,
                        tooltip="Editar",
                        on_click=lambda e, p=item: editar_produto(e, p)
                    ),

                    ft.IconButton(
                        icon=ft.Icons.DELETE,
                        tooltip="Excluir",
                        on_click=lambda e, p=item: excluir_produto(e, p)
                    )

                ])

            tabela_estoque.rows.append(

                ft.DataRow(

                    cells=[

                        ft.DataCell(ft.Text(status)),
                        ft.DataCell(ft.Text(item["nome"])),
                        ft.DataCell(ft.Text(item["unidade_medida"])),
                        ft.DataCell(ft.Text(str(item["quantidade"]))),
                        ft.DataCell(ft.Text(item["lote"])),
                        ft.DataCell(ft.Text(item["validade"])),

                        ft.DataCell(
                            ft.Row(
                                controls=acoes
                            )
                        )

                    ]

                )

            )

    def pagina_anterior(e):

        nonlocal pagina_atual

        if pagina_atual > 1:

            pagina_atual -= 1

            carregar_tabela(
                campo_pesquisa.value,
                filtro_dropdown.value
            )

            e.page.update()


    def proxima_pagina(e):

        nonlocal pagina_atual

        pagina_atual += 1

        carregar_tabela(
            campo_pesquisa.value,
            filtro_dropdown.value
        )

        e.page.update()


    botao_anterior = ft.Button(
        "Anterior",
        on_click=pagina_anterior
    )

    botao_proximo = ft.Button(
        "Próxima",
        on_click=proxima_pagina
    )
                  
    def pesquisar_produto(e):

        nonlocal pagina_atual

        pagina_atual = 1

        carregar_tabela(
            campo_pesquisa.value,
            filtro_dropdown.value
        )

        e.page.update()

    campo_pesquisa.on_change = pesquisar_produto   


    def atualizar_tela(e):

        atualizar_estoque()
        atualizar_cards()

        carregar_tabela(
            campo_pesquisa.value,
            filtro_dropdown.value
        )

        e.page.update()

    def logo_path():

        for path in LOGO_PATHS:

            if path.exists():

                return path

        return None 

    def data_hora_relatorio():

        return datetime.now().strftime(
            "%d/%m/%Y %H:%M:%S"
        )  

    def nome_usuario():

        if isinstance(usuario, dict):

            return (

                usuario.get("nome")

                or usuario.get("email")

                or "Usuário não identificado"

            )

        return "Usuário não identificado" 

    def aplicar_estilo_excel(ws):

        fill = PatternFill(

            "solid",

            fgColor="1F4E78"

        )

        fonte = Font(

            color="FFFFFF",

            bold=True

        )

        borda = Border(

            left=Side(style="thin"),

            right=Side(style="thin"),

            top=Side(style="thin"),

            bottom=Side(style="thin")

        )

        for cell in ws[9]:

            cell.fill = fill

            cell.font = fonte

            cell.alignment = Alignment(

                horizontal="center"

            )

            cell.border = borda

        for row in ws.iter_rows(min_row=10):

            for cell in row:

                cell.border = borda

                cell.alignment = Alignment(

                    vertical="center"

                )

        for i in range(1,7):

            ws.column_dimensions[

                get_column_letter(i)

            ].width = 22    

    def exportar_excel(e):


        arquivo = novo_excel("Estoque")

        wb = Workbook()

        ws = wb.active

        ws.title = "Estoque"

        logo = logo_path()

        if logo:

            try:

                img = ExcelImage(str(logo))

                img.width = 120

                img.height = 60

                ws.add_image(img, "A1")

            except:

                pass

        ws.merge_cells("B1:G1")

        ws["B1"] = "STOCK FLOW"

        ws["B1"].font = Font(

            size=18,

            bold=True

        )

        ws.merge_cells("B2:G2")

        ws["B2"] = "Relatório de Estoque"

        ws["B2"].font = Font(

            size=14,

            bold=True

        )

        ws["A4"] = "Gerado em:"

        ws["B4"] = data_hora_relatorio()

        ws["A5"] = "Gerado por:"

        ws["B5"] = nome_usuario()

        ws["A6"] = "Total itens:"

        ws["B6"] = len(estoque)

        ws.append([])

        ws.append(

            [

                "Status",

                "Produto",

                "Unidade",

                "Quantidade",

                "Lote",

                "Validade"

            ]

        )

        for item in estoque:

            status = "🟢"

            try:

                validade = datetime.strptime(

                    item["validade"],

                    "%d/%m/%Y"

                )

                if validade < datetime.now():

                    status = "🔴"

                elif (

                    item["quantidade"]

                    <= item["estoque_minimo"]

                ):

                    status = "🟡"

            except:

                pass

            ws.append(

                [

                    status,

                    item["nome"],

                    item["unidade_medida"],

                    item["quantidade"],

                    item["lote"],

                    item["validade"]

                ]

            )

        aplicar_estilo_excel(ws)

        wb.save(arquivo)

        page.launch_url(url_download(arquivo))
        
        e.page.show_dialog(
            ft.AlertDialog(
                modal=True,
                title=ft.Text("Excel gerado"),
                content=ft.Text(
                    f"Excel gerado com sucesso em:\n{arquivo}"
                ),
                actions=[
                    ft.TextButton(
                        "OK",
                        on_click=lambda e: e.page.pop_dialog()
                    )
                ]
            )
        )

        e.page.update()

    def gerar_pdf(e):


        arquivo = novo_pdf("Estoque")

        doc = SimpleDocTemplate(

            str(arquivo),

            pagesize=landscape(A4),

            rightMargin=1 * cm,

            leftMargin=1 * cm,

            topMargin=1 * cm,

            bottomMargin=1 * cm

        )

        estilos = getSampleStyleSheet()

        elementos = []

        logo = logo_path()

        cabecalho = []

        if logo:

            try:

                cabecalho.append(

                    Image(

                        str(logo),

                        width=3 * cm,

                        height=1.5 * cm

                    )

                )

            except:

                cabecalho.append("")

        else:

            cabecalho.append("")

        cabecalho.append(

            Paragraph(

                "<b>Stock Flow</b><br/>"

                "Relatório de Estoque<br/>"

                f"Gerado em: {data_hora_relatorio()}<br/>"

                f"Gerado por: {nome_usuario()}",

                estilos["Normal"]

            )

        )

        tabela_cabecalho = Table(

            [cabecalho],

            colWidths=[4 * cm, 22 * cm]

        )

        tabela_cabecalho.setStyle(

            TableStyle(

                [

                    (

                        "VALIGN",

                        (0,0),

                        (-1,-1),

                        "MIDDLE"

                    ),

                    (

                        "BOX",

                        (0,0),

                        (-1,-1),

                        0.5,

                        colors.lightgrey

                    ),

                    (

                        "BACKGROUND",

                        (0,0),

                        (-1,-1),

                        colors.whitesmoke

                    ),

                    (

                        "PADDING",

                        (0,0),

                        (-1,-1),

                        8

                    )

                ]

            )

        )

        elementos.append(tabela_cabecalho)

        elementos.append(

            Spacer(1, 15)

        )

        elementos.append(

            Paragraph(

                f"<b>Total de itens:</b> {len(estoque)}",

                estilos["Normal"]

            )

        )

        elementos.append(

            Spacer(1, 10)

        )

        dados = [

            [

                "Status",

                "Produto",

                "Unidade",

                "Quantidade",

                "Lote",

                "Validade"

            ]

        ]

        for item in estoque:

            status = "🟢"

            try:

                validade = datetime.strptime(

                    item["validade"],

                    "%d/%m/%Y"

                )

                if validade < datetime.now():

                    status = "🔴"

                elif (

                    item["quantidade"]

                    <= item["estoque_minimo"]

                ):

                    status = "🟡"

            except:

                pass

            dados.append(

                [

                    status,

                    item["nome"],

                    item["unidade_medida"],

                    str(item["quantidade"]),

                    item["lote"],

                    item["validade"]

                ]

            )

        tabela = Table(

            dados,

            repeatRows=1

        )

        tabela.setStyle(

            TableStyle(

                [

                    (

                        "BACKGROUND",

                        (0,0),

                        (-1,0),

                        colors.HexColor("#1F4E78")

                    ),

                    (

                        "TEXTCOLOR",

                        (0,0),

                        (-1,0),

                        colors.white

                    ),

                    (

                        "FONTNAME",

                        (0,0),

                        (-1,0),

                        "Helvetica-Bold"

                    ),

                    (

                        "GRID",

                        (0,0),

                        (-1,-1),

                        0.25,

                        colors.grey

                    ),

                    (

                        "FONTSIZE",

                        (0,0),

                        (-1,-1),

                        8

                    ),

                    (

                        "VALIGN",

                        (0,0),

                        (-1,-1),

                        "TOP"

                    ),

                    (

                        "ROWBACKGROUNDS",

                        (0,1),

                        (-1,-1),

                        [

                            colors.white,

                            colors.whitesmoke

                        ]

                    )

                ]

            )

        )

        elementos.append(tabela)

        elementos.append(

            Spacer(1, 15)

        )

        elementos.append(

            Paragraph(

                "<b>Stock Flow © 2026</b><br/>"

                "Sistema de Controle de Estoque",

                estilos["Normal"]

            )

        )

        try:


            doc.build(elementos)

            page.launch_url(url_download(arquivo))

            page.snack_bar = ft.SnackBar(
                content=ft.Text(
                    f"PDF gerado com sucesso:\n{arquivo}"
                )
            )

            page.snack_bar.open = True
            page.update()

        except Exception as erro:

            print("ERRO AO GERAR PDF:")
            print(erro)   
        
    carregar_tabela()                         

    return ft.Container(

        expand=True,

        bgcolor=BRANCO,

        alignment=ft.Alignment(0, 0),

        padding=15,

        content=ft.Container(
            width=min(largura_util, 1100),
            alignment=ft.Alignment(0, -1),
            content=ft.Column(
                [
            ft.Row(
                [
                    ft.Icon(
                        ft.Icons.INVENTORY,
                        size=34
                    ),

                    ft.Text(
                        "CONTROLE DE ESTOQUE",
                        size=24,
                        weight=ft.FontWeight.BOLD,
                        color=TEXTO
                    )
                ],
                wrap=True,
                spacing=10,
                run_spacing=10,
                alignment=ft.MainAxisAlignment.CENTER
            ),

            ft.Row(

                [

                    card_padrao(

                        content=ft.Container(

                            content=ft.Column(

                                [

                                    ft.Text(
                                        "Total",
                                        size=16,
                                        weight=ft.FontWeight.BOLD,
                                        color=TEXTO
                                    ),

                                    ft.Divider(

                                    color=VERDE,

                                    thickness=1

                                    ),

                                    texto_total_itens

                                ],

                                horizontal_alignment=ft.CrossAxisAlignment.CENTER

                            ),

                            padding=15,
                            width=largura_card

                        )

                    ),

                    card_padrao(

                        content=ft.Container(

                            content=ft.Column(

                                [

                                    ft.Text(
                                        "🟡 Estoque Baixo",
                                        size=16,
                                        weight=ft.FontWeight.BOLD,
                                        color=TEXTO
                                    ),

                                    ft.Divider(

                                    color=VERDE,

                                    thickness=1

                                    ),

                                    texto_total_baixo

                                ],

                                horizontal_alignment=ft.CrossAxisAlignment.CENTER

                            ),

                            padding=15,
                            width=largura_card

                        )

                    ),

                    card_padrao(

                        content=ft.Container(

                            content=ft.Column(

                                [

                                    ft.Text(
                                        "🔴 Vencidos",
                                        size=16,
                                        weight=ft.FontWeight.BOLD,
                                        color=TEXTO
                                    ),

                                    ft.Divider(

                                    color=VERDE,

                                    thickness=1

                                    ),

                                    texto_total_vencidos

                                ],

                                horizontal_alignment=ft.CrossAxisAlignment.CENTER

                            ),

                            padding=15,
                            width=largura_card

                        )

                    )

                ],
                wrap=True,
                spacing=10,
                run_spacing=10,
                alignment=ft.MainAxisAlignment.CENTER
            ),

            ft.Row(

                [

                    botao_primario(
                        "Exportar Excel",
                        icon=ft.Icons.DOWNLOAD,
                        on_click=exportar_excel
                    ),

                    botao_primario(
                        "Gerar PDF",
                        icon=ft.Icons.DOWNLOAD,
                        on_click=gerar_pdf
                    ),

                    botao_primario(
                        "Atualizar",
                        icon=ft.Icons.REFRESH,
                        on_click=atualizar_tela
                     ),

                ],
                wrap=True,
                spacing=10,
                run_spacing=10,
                alignment=ft.MainAxisAlignment.CENTER

            ),


            ft.Divider(
                color=VERDE,
                thickness=1
            ),

            ft.Row(
                [
                    campo_pesquisa,
                    filtro_dropdown,
                    ordenacao_dropdown
                ],
                wrap=True,
                spacing=10,
                run_spacing=10,
                alignment=ft.MainAxisAlignment.CENTER
            ),

            ft.Row(
                controls=[
                    tabela_estoque
                ],
                scroll=ft.ScrollMode.AUTO
            ),

            ft.Row(
                [
                    botao_anterior,
                    texto_pagina,
                    botao_proximo
                ],
                alignment=ft.MainAxisAlignment.CENTER
            )

            

            ],
            horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            scroll=ft.ScrollMode.AUTO,
            spacing=10
        )
    )
)