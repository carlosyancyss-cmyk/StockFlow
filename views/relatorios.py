import flet as ft
import webbrowser

from datetime import datetime
from pathlib import Path

from utils.export_manager import (
    novo_pdf,
    novo_excel,
    url_download
)
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app_paths import LOGO_PATHS
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    Image,
)

from views.tema import (
    AZUL,
    VERDE,
    BRANCO,
    FUNDO,
    TEXTO,
    CINZA,
    PADDING_CARD,
    ESPACAMENTO,
    card_padrao,
    botao_primario,
    botao_sucesso,
    botao_secundario,
    titulo_pagina,
    container_pagina
)

from database.relatorios_supabase import (
    listar_fornecedores_relatorio,
    listar_setores_relatorio,
    listar_produtos_relatorio,
    listar_ajustes_estoque,
    consultar_entradas_relatorio,
    consultar_saidas_relatorio,
    consultar_movimentacao_produto_relatorio,
    buscar_estoque_atual_produto,
    listar_origens_relatorio,
    listar_setores_saida_relatorio
)


# BASE_DIR = Path(__file__).resolve().parent.parent
# REPORTS_DIR = BASE_DIR / "reports"
# LOGO_PATHS = [
#    BASE_DIR / "assets" / "logo.png",
#    BASE_DIR / "assets" / "logo" / "logo_nova.png",
#    BASE_DIR / "logo.png",
#]


def relatorios_view(page, usuario):

    largura_tela = page.width or 390

    largura_util = max(largura_tela - 30, 300)

    largura_campo_pequeno = min(largura_util, 220)
    largura_campo_medio = min(largura_util, 260)
    largura_campo_grande = min(largura_util, 300)


    mensagem = ft.Text(
        size=16,
        weight=ft.FontWeight.BOLD
    )

    resultado_resumo = ft.Text(
        "Nenhum relatório gerado ainda.",
        size=14
    )

    tabela_resultado = ft.DataTable(

        columns=[

            ft.DataColumn(
                ft.Text("Relatório")
            )

        ],

        rows=[]

    )

    # =========================
    # FILTROS - ENTRADAS
    # =========================

    entrada_data_inicio = ft.TextField(
        label="Data inicial (dd/mm/aaaa)",
        width=largura_campo_pequeno
    )

    entrada_data_fim = ft.TextField(
        label="Data final (dd/mm/aaaa)",
        width=largura_campo_pequeno
    )

    entrada_tipo = ft.Dropdown(
        label="Tipo de entrada",
        width=largura_campo_medio,
        options=[
            ft.dropdown.Option("TODOS"),
            ft.dropdown.Option("Nota Fiscal"),
            ft.dropdown.Option("Empréstimo"),
            ft.dropdown.Option("Doação"),
            ft.dropdown.Option("Inventário"),
            ft.dropdown.Option("Ajuste de Estoque")
        ],
        value="TODOS"
    )

    entrada_fornecedor = ft.Dropdown(
        label="Fornecedor",
        width=largura_campo_grande,
        visible=False,
        options=[]
    )

    # =========================
    # FILTROS - SAÍDAS
    # =========================

    saida_data_inicio = ft.TextField(
        label="Data inicial (dd/mm/aaaa)",
        width=largura_campo_pequeno
    )

    saida_data_fim = ft.TextField(
        label="Data final (dd/mm/aaaa)",
        width=largura_campo_pequeno
    )

    saida_tipo = ft.Dropdown(
        label="Tipo de saída",
        width=largura_campo_medio,
        options=[
            ft.dropdown.Option("TODOS"),
            ft.dropdown.Option("Saída para Setor"),
            ft.dropdown.Option("Saída por Empréstimo"),
            ft.dropdown.Option("Saída por Doação"),
            ft.dropdown.Option("Saída por Avaria"),
            ft.dropdown.Option("Saída por Vencimento"),
            ft.dropdown.Option("Ajuste de estoque")
        ],
        value="TODOS"
    )

    saida_setor = ft.Dropdown(
        label="Setor",
        width=largura_campo_grande,
        visible=False,
        options=[]
    )

    ajuste_data_inicio = ft.TextField(
        label="Data inicial (dd/mm/aaaa)",
        width=largura_campo_pequeno
    )

    ajuste_data_fim = ft.TextField(
        label="Data final (dd/mm/aaaa)",
        width=largura_campo_pequeno
    )

    ajuste_usuario = ft.TextField(
        label="Usuário",
        width=largura_campo_grande
    )

    produto_relatorio = ft.Dropdown(
        label="Produto",
        width=largura_campo_grande,
        options=[]
    )

    produto_data_inicio = ft.TextField(
        label="Data Inicial (dd/mm/aaaa)",
        width=largura_campo_pequeno
    )

    produto_data_fim = ft.TextField(
        label="Data Final (dd/mm/aaaa)",
        width=largura_campo_pequeno
    )

    def formatar_data(e):

        texto = "".join(
            c for c in e.control.value
            if c.isdigit()
        )

        if len(texto) > 8:
            texto = texto[:8]

        if len(texto) >= 5:
            texto = f"{texto[:2]}/{texto[2:4]}/{texto[4:]}"
        elif len(texto) >= 3:
            texto = f"{texto[:2]}/{texto[2:]}"

        e.control.value = texto
        e.page.update()

    dados_atuais = []
    colunas_atuais = []
    titulo_atual = ""
    filtros_atuais = []

    entrada_data_inicio.on_change = formatar_data
    entrada_data_fim.on_change = formatar_data

    saida_data_inicio.on_change = formatar_data
    saida_data_fim.on_change = formatar_data

    ajuste_data_inicio.on_change = formatar_data
    ajuste_data_fim.on_change = formatar_data

    produto_data_inicio.on_change = formatar_data
    produto_data_fim.on_change = formatar_data

    def nome_usuario():

        if isinstance(usuario, dict):
            return (
                usuario.get("nome") or
                usuario.get("email") or
                usuario.get("login") or
                "Usuário não identificado"
            )

        return "Usuário não identificado"

    def logo_path():

        for path in LOGO_PATHS:
            if path.exists():
                return path

        return None

    def data_sql(valor):

        if not valor:
            return None

        try:
            return datetime.strptime(
                valor,
                "%d/%m/%Y"
            ).strftime("%Y-%m-%d")
        except:
            return None

    def data_hora_relatorio():
        return datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    def limpar_resultado():

        tabela_resultado.columns = [

            ft.DataColumn(
                ft.Text("Relatório")
            )

        ]

        tabela_resultado.rows.clear()

        resultado_resumo.value = (
            "Nenhum relatório gerado ainda."
        )

    def carregar_fornecedores():

        fornecedores = listar_fornecedores_relatorio()

        entrada_fornecedor.options = [
            ft.dropdown.Option("TODOS")
        ] + [
            ft.dropdown.Option(item["fornecedor"])
            for item in fornecedores
        ]

        entrada_fornecedor.value = "TODOS"

    def carregar_setores():

        setores = listar_setores_relatorio()

        saida_setor.options = [
            ft.dropdown.Option("TODOS")
        ] + [
            ft.dropdown.Option(item["setor"])
            for item in setores
        ]

        saida_setor.value = "TODOS"

    def carregar_produtos_relatorio():

        produtos = listar_produtos_relatorio()

        produto_relatorio.options = [

            ft.dropdown.Option(item["nome"])

            for item in produtos

        ]    

    def alterar_tipo_entrada(e):

        if entrada_tipo.value == "Nota Fiscal":
            entrada_fornecedor.label = "Fornecedor"
            carregar_fornecedores()
            entrada_fornecedor.visible = True

        elif entrada_tipo.value in ["Empréstimo", "Doação"]:
            entrada_fornecedor.label = "Instituição"
            carregar_instituicoes_entrada(entrada_tipo.value)
            entrada_fornecedor.visible = True

        else:
            entrada_fornecedor.visible = False
            entrada_fornecedor.value = "TODOS"

        e.page.update()

    def alterar_tipo_saida(e):

        if saida_tipo.value == "Saída para Setor":
            saida_setor.label = "Setor"
            carregar_setores()
            saida_setor.visible = True

        elif saida_tipo.value in ["Saída por Empréstimo", "Saída por Doação"]:
            saida_setor.label = "Instituição"
            carregar_instituicoes_saida(saida_tipo.value)
            saida_setor.visible = True

        else:
            saida_setor.visible = False
            saida_setor.value = "TODOS"

        e.page.update()

    entrada_tipo.on_select = alterar_tipo_entrada
    saida_tipo.on_select = alterar_tipo_saida

    carregar_fornecedores()
    carregar_setores()
    carregar_produtos_relatorio()

    def carregar_instituicoes_entrada(tipo):

        origens = listar_origens_relatorio(tipo)

        entrada_fornecedor.options = [
            ft.dropdown.Option("TODOS")
        ] + [
            ft.dropdown.Option(origem)
            for origem in origens
        ]

        entrada_fornecedor.value = "TODOS"


    def carregar_instituicoes_saida(tipo):

        setores = listar_setores_saida_relatorio(tipo)

        saida_setor.options = [
            ft.dropdown.Option("TODOS")
        ] + [
            ft.dropdown.Option(setor)
            for setor in setores
        ]

        saida_setor.value = "TODOS"

    def montar_tabela(colunas, dados):

        tabela_resultado.columns.clear()
        tabela_resultado.rows.clear()

        if not colunas:

            tabela_resultado.columns = [

                ft.DataColumn(
                    ft.Text("Relatório")
                )

            ]

            return

        tabela_resultado.columns.extend(
            [
                ft.DataColumn(ft.Text(coluna))
                for coluna in colunas
            ]
        )

        for linha in dados[:100]:

            tabela_resultado.rows.append(
                ft.DataRow(
                    cells=[
                        ft.DataCell(
                            ft.Text(str(linha.get(coluna, "")))
                        )
                        for coluna in colunas
                    ]
                )
            )

    def consultar_entradas():

        return consultar_entradas_relatorio()

    def consultar_saidas():

        return consultar_saidas_relatorio()

    def filtros_entradas():

        filtros = [
            f"Data inicial: {entrada_data_inicio.value or 'Todas'}",
            f"Data final: {entrada_data_fim.value or 'Todas'}",
            f"Tipo: {entrada_tipo.value or 'TODOS'}"
        ]

        if entrada_tipo.value == "Nota Fiscal":
            filtros.append(
                f"Fornecedor: {entrada_fornecedor.value or 'TODOS'}"
            )

        return filtros

    def filtros_saidas():

        filtros = [
            f"Data inicial: {saida_data_inicio.value or 'Todas'}",
            f"Data final: {saida_data_fim.value or 'Todas'}",
            f"Tipo: {saida_tipo.value or 'TODOS'}"
        ]

        if saida_tipo.value == "Saída para Setor":
            filtros.append(
                f"Setor: {saida_setor.value or 'TODOS'}"
            )

        return filtros
    
    def consultar_movimentacao_produto():

        if not produto_relatorio.value:
            return []

        return consultar_movimentacao_produto_relatorio(
            produto_relatorio.value
        )

    def gerar_relatorio_entradas(e):

        nonlocal dados_atuais
        nonlocal colunas_atuais
        nonlocal titulo_atual
        nonlocal filtros_atuais

        titulo_atual = "Relatório de Entradas"
        filtros_atuais = filtros_entradas()
        dados_atuais = consultar_entradas()
        # Data inicial
        if entrada_data_inicio.value:
            inicio = data_sql(entrada_data_inicio.value)
            dados_atuais = [
                x for x in dados_atuais
                if x["data"][:10] >= inicio
            ]

        # Data final
        if entrada_data_fim.value:
            fim = data_sql(entrada_data_fim.value)
            dados_atuais = [
                x for x in dados_atuais
                if x["data"][:10] <= fim
            ]

        # Tipo
        if entrada_tipo.value and entrada_tipo.value != "TODOS":
            dados_atuais = [
                x for x in dados_atuais
                if x["tipo"] == entrada_tipo.value
            ]

        # Fornecedor / Origem
        if (
            entrada_fornecedor.value
            and entrada_fornecedor.value != "TODOS"
        ):
            dados_atuais = [
                x for x in dados_atuais
                if (
                    x["fornecedor"] == entrada_fornecedor.value
                    or
                    x["origem"] == entrada_fornecedor.value
                )
            ]
        colunas_atuais = [
            "codigo",
            "tipo",
            "fornecedor",
            "origem",
            "numero_nf",
            "valor_total",
            "usuario",
            "data",
            "produto",
            "quantidade",
            "lote",
            "validade",
            "valor_unitario"
        ]

        montar_tabela(colunas_atuais, dados_atuais)

        resultado_resumo.value = (
            f"{len(dados_atuais)} registro(s) encontrado(s). "
            "A prévia mostra até 100 linhas."
        )

        mensagem.value = "Relatório de entradas gerado na tela."

        e.page.update()

    def gerar_relatorio_saidas(e):

        nonlocal dados_atuais
        nonlocal colunas_atuais
        nonlocal titulo_atual
        nonlocal filtros_atuais

        titulo_atual = "Relatório de Saídas"
        filtros_atuais = filtros_saidas()
        dados_atuais = consultar_saidas()
        # Data inicial
        if saida_data_inicio.value:
            inicio = data_sql(saida_data_inicio.value)
            dados_atuais = [
                x for x in dados_atuais
                if x["data"][:10] >= inicio
            ]

        # Data final
        if saida_data_fim.value:
            fim = data_sql(saida_data_fim.value)
            dados_atuais = [
                x for x in dados_atuais
                if x["data"][:10] <= fim
            ]

        # Tipo
        if saida_tipo.value and saida_tipo.value != "TODOS":
            dados_atuais = [
                x for x in dados_atuais
                if x["tipo"] == saida_tipo.value
            ]

        # Setor
        if (
            saida_setor.value
            and saida_setor.value != "TODOS"
        ):
            dados_atuais = [
                x for x in dados_atuais
                if x["setor"] == saida_setor.value
            ]
        colunas_atuais = [
            "codigo",
            "tipo",
            "setor",
            "motivo",
            "usuario",
            "data",
            "produto",
            "quantidade",
            "lote"
        ]

        montar_tabela(colunas_atuais, dados_atuais)

        resultado_resumo.value = (
            f"{len(dados_atuais)} registro(s) encontrado(s). "
            "A prévia mostra até 100 linhas."
        )

        mensagem.value = "Relatório de saídas gerado na tela."

        e.page.update()

    def gerar_relatorio_ajustes(e):

        nonlocal dados_atuais
        nonlocal colunas_atuais
        nonlocal titulo_atual
        nonlocal filtros_atuais

        inicio = data_sql(ajuste_data_inicio.value)
        fim = data_sql(ajuste_data_fim.value)

        titulo_atual = "Relatório de Ajustes de Estoque"

        filtros_atuais = [
            f"Data inicial: {ajuste_data_inicio.value or 'Todas'}",
            f"Data final: {ajuste_data_fim.value or 'Todas'}",
            f"Usuário: {ajuste_usuario.value or 'Todos'}"
        ]

        dados_atuais = [
            dict(item)
            for item in listar_ajustes_estoque(
                inicio,
                fim,
                ajuste_usuario.value
            )
        ]

        colunas_atuais = [
            "data",
            "usuario",
            "produto",
            "acao",
            "antes",
            "depois"
        ]

        montar_tabela(colunas_atuais, dados_atuais)

        resultado_resumo.value = (
            f"{len(dados_atuais)} ajuste(s) encontrado(s). "
            "A prévia mostra até 100 linhas."
        )

        mensagem.value = "Relatório de ajustes de estoque gerado na tela."

        e.page.update() 

    def gerar_relatorio_produto(e):

        nonlocal dados_atuais
        nonlocal colunas_atuais
        nonlocal titulo_atual
        nonlocal filtros_atuais

        if not produto_relatorio.value:

            mensagem.value = "Selecione um produto."

            e.page.update()

            return

        titulo_atual = (
            f"Movimentação do Produto - "
            f"{produto_relatorio.value}"
        )

        filtros_atuais = [

            f"Produto: {produto_relatorio.value}",

            f"Data Inicial: {produto_data_inicio.value or 'Todas'}",

            f"Data Final: {produto_data_fim.value or 'Todas'}"

        ]

        dados_atuais = consultar_movimentacao_produto()

        print("\n===== MOVIMENTAÇÕES =====")

        for item in dados_atuais:
            print(
                item["data_mov"],
                item["movimento"],
                item["quantidade"],
                item.get("antes"),
                item.get("depois")
            )

        print("=========================\n")

        estoque_atual = buscar_estoque_atual_produto(
            produto_relatorio.value
        )

        saldo = 0

        total_entradas = 0
        total_saidas = 0
        total_ajustes = 0

        for item in dados_atuais:

            qtd = int(item["quantidade"] or 0)

            if item["movimento"] == "ENTRADA":

                saldo += qtd
                total_entradas += qtd

            elif item["movimento"] == "SAIDA":

                saldo -= qtd
                total_saidas += qtd

            elif item["movimento"] == "AJUSTE":

                # Ajustes são apenas eventos de auditoria.
                # Não alteram o saldo acumulado do relatório.

                total_ajustes += 1

            item["saldo"] = saldo

        colunas_atuais = [

            "data_mov",
            "movimento",
            "tipo",
            "produto",
            "quantidade",
            "saldo",
            "lote",
            "usuario"

        ]

        montar_tabela(
            colunas_atuais,
            dados_atuais
        )

        resultado_resumo.value = (

            f"Produto: {produto_relatorio.value}\n\n"

            f"Total Entradas: {total_entradas}\n"

            f"Total Saídas: {total_saidas}\n"

            f"Total Ajustes: {total_ajustes}\n"

            f"Saldo Atual: {estoque_atual}\n\n"

            f"Movimentações encontradas: {len(dados_atuais)}"

        )

        mensagem.value = (

            "Relatório de movimentação do produto gerado."

        )

        e.page.update()       

    def validar_relatorio():

        if not dados_atuais:
            mensagem.value = (
                "Gere um relatório antes de exportar."
            )
            return False

        return True

    def nome_arquivo(base, extensao):

        if extensao == "pdf":
            return novo_pdf(base)

        if extensao == "xlsx":
            return novo_excel(base)
        
    def aplicar_estilo_excel(ws, total_colunas):

        fill_header = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        fonte_header = Font(
            color="FFFFFF",
            bold=True
        )

        borda = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        for cell in ws[8]:
            cell.fill = fill_header
            cell.font = fonte_header
            cell.alignment = Alignment(horizontal="center")
            cell.border = borda

        for row in ws.iter_rows(min_row=9):
            for cell in row:
                cell.border = borda
                cell.alignment = Alignment(vertical="top")

        for col_idx in range(1, total_colunas + 1):
            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = 20

    async def exportar_excel(e):

        if not validar_relatorio():
            e.page.update()
            return

        arquivo = nome_arquivo("Relatorio", "xlsx")

        wb = Workbook()
        ws = wb.active
        ws.title = "Relatório"

        total_colunas = len(colunas_atuais)

        logo = logo_path()

        if logo:
            try:
                imagem = ExcelImage(str(logo))
                imagem.width = 120
                imagem.height = 60
                ws.add_image(imagem, "A1")
            except:
                pass

        ws.merge_cells(
            start_row=1,
            start_column=2,
            end_row=1,
            end_column=total_colunas
        )
        ws.cell(1, 2).value = "Stock Flow"
        ws.cell(1, 2).font = Font(size=18, bold=True)

        ws.merge_cells(
            start_row=2,
            start_column=2,
            end_row=2,
            end_column=total_colunas
        )
        ws.cell(2, 2).value = titulo_atual
        ws.cell(2, 2).font = Font(size=14, bold=True)

        ws.cell(4, 1).value = "Gerado em:"
        ws.cell(4, 2).value = data_hora_relatorio()

        ws.cell(5, 1).value = "Gerado por:"
        ws.cell(5, 2).value = nome_usuario()

        ws.cell(6, 1).value = "Filtros:"
        ws.cell(6, 2).value = " | ".join(filtros_atuais)

        ws.append([])
        ws.append(colunas_atuais)

        for item in dados_atuais:
            ws.append(
                [
                    item.get(coluna, "")
                    for coluna in colunas_atuais
                ]
            )

        aplicar_estilo_excel(ws, total_colunas)

        wb.save(arquivo)

        await page.launch_url(url_download(arquivo))

        mensagem.value = f"Excel gerado: {arquivo}"

        e.page.update()

    async def exportar_pdf(e):

        print("ENTROU EM exportar_pdf")

        if not validar_relatorio():
            e.page.update()
            return

        arquivo = nome_arquivo("Relatorio", "pdf")

        doc = SimpleDocTemplate(
            str(arquivo),
            pagesize=landscape(A4),
            rightMargin=1 * cm,
            leftMargin=1 * cm,
            topMargin=1 * cm,
            bottomMargin=1 * cm
        )

        estilos = getSampleStyleSheet()
        elementos = []

        logo = logo_path()

        cabecalho = []

        if logo:
            try:
                cabecalho.append(
                    Image(str(logo), width=3.0 * cm, height=1.5 * cm)
                )
            except:
                cabecalho.append("")
        else:
            cabecalho.append("")

        cabecalho.append(
            Paragraph(
                "<b>Stock Flow</b><br/>"
                f"{titulo_atual}<br/>"
                f"Gerado em: {data_hora_relatorio()}<br/>"
                f"Gerado por: {nome_usuario()}",
                estilos["Normal"]
            )
        )

        tabela_cabecalho = Table(
            [cabecalho],
            colWidths=[4 * cm, 22 * cm]
        )

        tabela_cabecalho.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.whitesmoke),
                    ("PADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )

        elementos.append(tabela_cabecalho)
        elementos.append(Spacer(1, 10))

        elementos.append(
            Paragraph(
                "<b>Filtros:</b> " + " | ".join(filtros_atuais),
                estilos["Normal"]
            )
        )

        elementos.append(Spacer(1, 10))

        dados_tabela = [
            colunas_atuais
        ]

        estilo_celula = estilos["BodyText"]
        estilo_celula.fontName = "Helvetica"
        estilo_celula.fontSize = 7
        estilo_celula.leading = 9

        for item in dados_atuais:

            linha = []

            for coluna in colunas_atuais:

                valor = str(item.get(coluna, ""))

                linha.append(
                    Paragraph(
                        valor,
                        estilo_celula
                    )
                )

            dados_tabela.append(linha)

        if titulo_atual == "Relatório de Ajustes de Estoque":

            larguras = [
                2.5 * cm,   # data
                3.0 * cm,   # usuario
                5.5 * cm,   # produto
                3.0 * cm,   # acao
                6.0 * cm,   # antes
                6.0 * cm    # depois
            ]

            tabela = Table(
                dados_tabela,
                colWidths=larguras,
                repeatRows=1
            )

        else:

            tabela = Table(
                dados_tabela,
                repeatRows=1
            )

        tabela.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.whitesmoke]),
                ]
            )
        )

        elementos.append(tabela)

        print("GERANDO PDF...")

        doc.build(elementos)

        print("PDF SALVO EM:")
        print(arquivo)
        print("EXISTE APÓS GERAR?", Path(arquivo).exists())

        await page.launch_url(url_download(arquivo))

        mensagem.value = "PDF gerado com sucesso."

        page.snack_bar = ft.SnackBar(
            content=ft.Text(
                f"PDF salvo em:\n{arquivo}"
            )
        )

        page.snack_bar.open = True

        e.page.update()

    conteudo_relatorio = ft.Column()

    def mostrar_relatorio_entradas(e=None):

        conteudo_relatorio.controls.clear()

        conteudo_relatorio.controls.append(
            ft.Container(
                padding=15,
                content=ft.Column(
                    [
                        ft.Text(
                            "Relatório de Entradas",
                            size=20,
                            weight=ft.FontWeight.BOLD,
                            color=TEXTO
                        ),

                        ft.Row(
                            [
                                entrada_data_inicio,
                                entrada_data_fim,
                                entrada_tipo,
                                entrada_fornecedor
                            ],
                            wrap=True,
                            spacing=10,
                            run_spacing=10,
                            alignment=ft.MainAxisAlignment.CENTER
                        ),

                        ft.Row(
                            [
                                botao_primario(
                                    "Gerar Relatório",
                                    icon=ft.Icons.SEARCH,
                                    on_click=gerar_relatorio_entradas
                                ),

                                botao_primario(
                                    "Exportar Excel",
                                    icon=ft.Icons.DOWNLOAD,
                                    on_click=exportar_excel
                                ),

                                botao_primario(
                                    "Exportar PDF",
                                    icon=ft.Icons.PICTURE_AS_PDF,
                                    on_click=exportar_pdf
                                )
                            ],
                            wrap=True,
                            spacing=10,
                            run_spacing=10,
                            alignment=ft.MainAxisAlignment.CENTER
                        )
                    ]
                )
            )
        )

        if e:
            e.page.update()


    def mostrar_relatorio_saidas(e=None):

        conteudo_relatorio.controls.clear()

        conteudo_relatorio.controls.append(
            ft.Container(
                padding=15,
                content=ft.Column(
                    [
                        ft.Text(
                            "Relatório de Saídas",
                            size=20,
                            weight=ft.FontWeight.BOLD,
                            color=TEXTO
                        ),

                        ft.Row(
                            [
                                saida_data_inicio,
                                saida_data_fim,
                                saida_tipo,
                                saida_setor
                            ],
                            wrap=True,
                            spacing=10,
                            run_spacing=10,
                            alignment=ft.MainAxisAlignment.CENTER
                        ),

                        ft.Row(
                            [
                                botao_primario(
                                    "Gerar Relatório",
                                    icon=ft.Icons.SEARCH,
                                    on_click=gerar_relatorio_saidas
                                ),

                                botao_primario(
                                    "Exportar Excel",
                                    icon=ft.Icons.DOWNLOAD,
                                    on_click=exportar_excel
                                ),

                                botao_primario(
                                    "Exportar PDF",
                                    icon=ft.Icons.PICTURE_AS_PDF,
                                    on_click=exportar_pdf
                                )
                            ],
                            wrap=True,
                            spacing=10,
                            run_spacing=10,
                            alignment=ft.MainAxisAlignment.CENTER
                        )
                    ]
                )
            )
        )

        if e:
            e.page.update()

    def mostrar_relatorio_ajustes(e=None):

        conteudo_relatorio.controls.clear()

        conteudo_relatorio.controls.append(
            ft.Container(
                padding=15,
                content=ft.Column(
                    [
                        ft.Text(
                            "Relatório de Ajustes de Estoque",
                            size=20,
                            weight=ft.FontWeight.BOLD,
                            color=TEXTO
                        ),

                        ft.Row(
                            [
                                ajuste_data_inicio,
                                ajuste_data_fim,
                                ajuste_usuario
                            ],
                            wrap=True,
                            spacing=10,
                            run_spacing=10,
                            alignment=ft.MainAxisAlignment.CENTER
                        ),

                        ft.Row(
                            [
                                botao_primario(
                                    "Gerar Relatório",
                                    icon=ft.Icons.SEARCH,
                                    on_click=gerar_relatorio_ajustes
                                ),

                                botao_primario(
                                    "Exportar Excel",
                                    icon=ft.Icons.DOWNLOAD,
                                    on_click=exportar_excel
                                ),

                                botao_primario(
                                    "Exportar PDF",
                                    icon=ft.Icons.PICTURE_AS_PDF,
                                    on_click=exportar_pdf
                                )
                            ],
                            wrap=True,
                            spacing=10,
                            run_spacing=10,
                            alignment=ft.MainAxisAlignment.CENTER
                        )
                    ]
                )
            )
        )

        if e:
            e.page.update()

    def mostrar_relatorio_produto(e=None):

        conteudo_relatorio.controls.clear()

        conteudo_relatorio.controls.append(

            ft.Container(

                padding=15,

                content=ft.Column(

                    [

                        ft.Text(

                            "Movimentação por Produto",

                            size=20,

                            weight=ft.FontWeight.BOLD,

                            color=TEXTO

                        ),

                        ft.Row(

                            [

                                produto_relatorio,

                                produto_data_inicio,

                                produto_data_fim

                            ],

                            wrap=True,

                            spacing=10,

                            alignment=ft.MainAxisAlignment.CENTER

                        ),

                        ft.Row(

                            [

                                botao_primario(

                                    "Gerar Relatório",

                                    icon=ft.Icons.SEARCH,

                                    on_click=gerar_relatorio_produto

                                ),

                                botao_primario(

                                    "Exportar Excel",

                                    icon=ft.Icons.DOWNLOAD,

                                    on_click=exportar_excel

                                ),

                                botao_primario(

                                    "Exportar PDF",

                                    icon=ft.Icons.PICTURE_AS_PDF,

                                    on_click=exportar_pdf

                                )

                            ],

                            wrap=True,

                            spacing=10,

                            alignment=ft.MainAxisAlignment.CENTER

                        )

                    ]

                )

            )

        )

        if e:

            e.page.update()                


    mostrar_relatorio_entradas()    

    return ft.Container(

        expand=True,

        bgcolor=BRANCO,

        alignment=ft.Alignment(0, 0),

        padding=15,

        content=ft.Container(
            width=min(largura_util, 1100),
            alignment=ft.Alignment(0, -1),
            content=ft.Column(
                [
            ft.Row(
                [
                    ft.Icon(
                        ft.Icons.ASSESSMENT,
                        size=34
                    ),
                    ft.Text(
                        "RELATÓRIOS",
                        size=24,
                        weight=ft.FontWeight.BOLD,
                        color=TEXTO
                    )
                ],
                wrap=True,
                spacing=10,
                run_spacing=10,
                alignment=ft.MainAxisAlignment.CENTER
            ),

            ft.Text(
                "Todos os relatórios são exportados com logo, nome do sistema, data, hora, usuário e filtros aplicados.",
                size=14,
                color=TEXTO
            ),

            ft.Divider(
                color=VERDE,
                thickness=1
            ),

            ft.Row(
                [
                    botao_primario(
                        "Entradas",
                        icon=ft.Icons.INPUT,
                        on_click=mostrar_relatorio_entradas
                    ),

                    botao_primario(
                        "Saídas",
                        icon=ft.Icons.OUTPUT,
                        on_click=mostrar_relatorio_saidas
                    ),

                    botao_primario(
                        "Ajustes Estoque",
                        icon=ft.Icons.EDIT_NOTE,
                        on_click=mostrar_relatorio_ajustes
                    ),

                    botao_primario(
                        "Por Produto",
                        icon=ft.Icons.INVENTORY_2,
                        on_click=mostrar_relatorio_produto
                    )

                ],
                wrap=True,
                spacing=10,
                run_spacing=10,
                alignment=ft.MainAxisAlignment.CENTER
            ),

            conteudo_relatorio,

            ft.Divider(
                color=VERDE,
                thickness=1
            ),

            resultado_resumo,

            ft.Row(
                controls=[
                    tabela_resultado
                ],
                scroll=ft.ScrollMode.AUTO
            ),

            mensagem
            ],
            horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            scroll=ft.ScrollMode.AUTO,
            spacing=10
        )
    )
)
